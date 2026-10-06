"""Read-only XLSX checker with a deliberately small, fail-closed evaluator.

Supports arithmetic, cell/range references, SUM and AVERAGE. No eval, macros,
external links, named formulas or network. Business totals require a contract.
"""
import ast
import operator
import re
from decimal import Decimal, DivisionByZero, InvalidOperation
from .gate import GateResult

CELL=re.compile(r"(?:(?:'([^']+)'|([A-Za-z_][\w ]*))!)?(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?(?![\w])")

class BlankReference(ValueError):pass
class Unsupported(ValueError):pass

def validate_xlsx(path,contract=None):
    try:
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
    except ImportError as e:
        raise ValueError('Excel检查需要可选依赖：pip install .[excel]') from e
    g=GateResult([])
    wb=openpyxl.load_workbook(path,read_only=True,data_only=False,keep_links=False)
    cached=openpyxl.load_workbook(path,read_only=True,data_only=True,keep_links=False)
    values={};memo={};active=set()
    try:
        for sheet in wb:
            if sheet.max_row*sheet.max_column>100000:
                g.add('Q030','BLOCK',sheet.title,'不超过100000个单元格','超出检查上限');continue
            for row in sheet:
                for cell in row:
                    if cell.value is not None:values[sheet.title,cell.coordinate]=cell.value
        def cell_value(sheet,coord):
            key=(sheet,coord.replace('$',''))
            if key in memo:return memo[key]
            if key in active:raise Unsupported('循环引用')
            if sheet not in wb.sheetnames:raise BlankReference('不存在的工作表 '+sheet)
            value=values.get(key)
            if value is None:raise BlankReference(f'{sheet}!{coord}为空（本工具不默认为0）')
            if isinstance(value,bool):raise Unsupported('不隐式将布尔值转换为金额')
            if isinstance(value,(int,float)):return Decimal(str(value))
            if not isinstance(value,str) or not value.startswith('='):raise InvalidOperation('非数字引用或错误值')
            active.add(key)
            try:
                formula=value[1:]
                if any(c in formula for c in ('[',']','"',';')):raise Unsupported('外部链接或不支持的公式语法')
                formula=CELL.sub(lambda m: f"RANGE({(m.group(1) or m.group(2) or sheet)!r},{m.group(3).replace('$','')!r},{m.group(4).replace('$','')!r})" if m.group(4) else f"CELL({(m.group(1) or m.group(2) or sheet)!r},{m.group(3).replace('$','')!r})",formula)
                formula=re.sub(r'(\d+(?:\.\d+)?)%',r'(\1/100)',formula)
                def evaluate(node):
                    if isinstance(node,ast.Constant):
                        if isinstance(node.value,(int,float)) and not isinstance(node.value,bool):return Decimal(str(node.value))
                        if isinstance(node.value,str):return node.value
                    if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.UAdd,ast.USub)):
                        v=evaluate(node.operand);return v if isinstance(node.op,ast.UAdd) else -v
                    if isinstance(node,ast.BinOp):
                        ops={ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul,ast.Div:operator.truediv}
                        if type(node.op) not in ops:raise Unsupported('不支持的运算符')
                        return ops[type(node.op)](evaluate(node.left),evaluate(node.right))
                    if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and not node.keywords:
                        args=[evaluate(n) for n in node.args];name=node.func.id.upper()
                        if name=='CELL' and len(args)==2:return cell_value(*args)
                        if name=='RANGE' and len(args)==3:
                            col1,row1,col2,row2=range_boundaries(args[1]+':'+args[2])
                            if (col2-col1+1)*(row2-row1+1)>10000:raise Unsupported('引用范围过大')
                            return [cell_value(args[0],c.coordinate) for row in wb[args[0]].iter_rows(min_row=row1,max_row=row2,min_col=col1,max_col=col2) for c in row]
                        flat=[v for arg in args for v in (arg if isinstance(arg,list) else [arg])]
                        if name=='SUM':return sum(flat,Decimal(0))
                        if name=='AVERAGE' and flat:return sum(flat,Decimal(0))/len(flat)
                    raise Unsupported('仅支持算术、CELL/RANGE、SUM、AVERAGE')
                answer=evaluate(ast.parse(formula,mode='eval').body)
                if not isinstance(answer,Decimal) or not answer.is_finite():raise Unsupported('公式未返回有限数字')
                memo[key]=answer;return answer
            finally:active.remove(key)
        for (sheet,coord),value in values.items():
            src=f'{path.name}!{sheet}!{coord}'
            if isinstance(value,str) and re.search(r'#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A|#NUM!|#NULL!',value):
                g.add('Q002','BLOCK',src,'无错误值',value);continue
            if isinstance(value,str) and value.startswith('='):
                try:
                    answer=cell_value(sheet,coord);cache=cached[sheet][coord].value
                    if isinstance(cache,(int,float)) and abs(answer-Decimal(str(cache)))>Decimal('.005'):
                        g.add('Q004','BLOCK',src,str(answer),cache,'公式缓存与本次复算不一致')
                    if isinstance(cache,str) and cache.startswith('#'):g.add('Q002','BLOCK',src,'无缓存错误',cache)
                except BlankReference as e:g.add('Q003','BLOCK',src,'有效且非空的数字引用',str(e))
                except (DivisionByZero,ZeroDivisionError,InvalidOperation,TypeError) as e:g.add('Q002','BLOCK',src,'可计算的数字公式',str(e))
                except (Unsupported,SyntaxError,ValueError,KeyError,RecursionError) as e:g.add('Q030','BLOCK',src,'在已支持的公式范围内',str(e),'不能验证的公式阻止发布，不能视为PASS')
        if contract is None:
            g.add('Q031','WARNING',path.name,'提供totals业务契约','未提供','已检查可解析公式；无法自动识别任意硬编码业务摘要')
        else:
            try:
                if set(contract)!={'totals'}:raise ValueError('契约仅支持totals字段')
                for c in contract['totals']:
                    if set(c)!={'sheet','components','total'}:raise ValueError('每项要求sheet/components/total')
                    expected=sum((cell_value(c['sheet'],x) for x in c['components']),Decimal(0))
                    actual=cell_value(c['sheet'],c['total'])
                    if expected!=actual:g.add('Q005','BLOCK',f"{path.name}!{c['sheet']}!{c['total']}",str(expected),str(actual))
            except (KeyError,TypeError,ValueError,InvalidOperation,DivisionByZero) as e:
                g.add('Q031','BLOCK',path.name,'可求值的检查契约',str(e))
        if not g.findings:g.add('Q030','PASS',path.name,'支持公式及契约通过','通过',action='无')
        return g
    finally:
        wb.close();cached.close()
